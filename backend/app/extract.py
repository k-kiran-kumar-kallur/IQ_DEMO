"""Turn each file into sections of text (for chunking) and tables (for SQLite)."""
import csv
import datetime as dt
import json
import re
from pathlib import Path

import docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
import openpyxl
import pptx
from pypdf import PdfReader

ROWS_PER_SECTION = 15


def _cell(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        return round(v, 6)
    return v


def _row_text(headers, row):
    parts = []
    for h, v in zip(headers, row):
        if v in (None, ""):
            continue
        if isinstance(v, float):
            v = f"{v:,.4f}".rstrip("0").rstrip(".") if abs(v) < 1 else f"{v:,.2f}".rstrip("0").rstrip(".")
        parts.append(f"{h}: {v}")
    return "; ".join(parts)


def _table_sections(label, headers, rows):
    out = []
    for i in range(0, len(rows), ROWS_PER_SECTION):
        block = rows[i:i + ROWS_PER_SECTION]
        text = "\n".join(_row_text(headers, r) for r in block)
        out.append({"label": f"{label} rows {i + 1}-{i + len(block)}", "text": text})
    return out


def _find_table(values):
    """Header = first row with 3+ cells that are all text, followed by a data row. Data stops at a blank first column."""
    for i, row in enumerate(values[:20]):
        cells = [c for c in row if c not in (None, "")]
        if len(cells) >= 3 and all(isinstance(c, str) for c in cells) and i + 1 < len(values):
            nxt = [c for c in values[i + 1] if c not in (None, "")]
            if len(nxt) >= 2:
                width = max(j for j, c in enumerate(row) if c not in (None, "")) + 1
                headers = [str(c) if c not in (None, "") else f"col{j + 1}" for j, c in enumerate(row[:width])]
                data = []
                for r in values[i + 1:]:
                    if not r or r[0] in (None, ""):
                        break
                    data.append([_cell(c) for c in r[:width]])
                return i, headers, data
    return None


def classification(path: Path, text: str, props_category=None):
    if props_category:
        return props_category
    m = re.search(r"\b(RESTRICTED|CONFIDENTIAL|INTERNAL|PUBLIC)\b", text[:3000])
    if m:
        return m.group(1).capitalize()
    rel = str(path)
    if "SysHCM" in rel:
        return "Restricted"
    if re.search(r"email|Email|Transcript|transcript|pasted|INC-|SLA_Breach|funding|contacts|Pulse", rel):
        return "Confidential"
    return "Internal"


def extract(path: Path):
    ext = path.suffix.lower().lstrip(".")
    title = path.stem.replace("_", " ")
    sections, tables, category, preface = [], [], None, ""

    if ext == "docx":
        d = docx.Document(str(path))
        category = d.core_properties.category or None
        title = d.core_properties.title or title
        current, buf = "Introduction", []

        def flush():
            if buf:
                sections.append({"label": current, "text": "\n".join(buf)})
        body = d.element.body
        for child in body.iterchildren():
            tag = child.tag.split("}")[-1]
            if tag == "p":
                p = DocxParagraph(child, d)
                t = p.text.strip()
                if not t:
                    continue
                if p.style.name.startswith("Heading"):
                    flush()
                    current, buf = t, []
                else:
                    buf.append(("- " if "List" in p.style.name else "") + t)
            elif tag == "tbl":
                tbl = DocxTable(child, d)
                rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
                if rows:
                    hdr = rows[0]
                    if len(rows) > 1 and any(hdr):
                        buf.extend(_row_text(hdr, r) for r in rows[1:])
                    else:
                        buf.extend(" | ".join(r) for r in rows)
        flush()

    elif ext == "pptx":
        p = pptx.Presentation(str(path))
        category = p.core_properties.category or None
        title = p.core_properties.title or title
        for i, s in enumerate(p.slides, 1):
            texts, slide_title = [], ""
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    t = sh.text_frame.text.strip()
                    if sh.name == "Title":
                        slide_title = t
                    texts.append(t)
                if sh.has_table:
                    rows = [[c.text for c in r.cells] for r in sh.table.rows]
                    texts.extend(_row_text(rows[0], r) for r in rows[1:])
                if sh.has_chart:
                    ch = sh.chart
                    try:
                        cats = list(ch.plots[0].categories)
                        for ser in ch.plots[0].series:
                            texts.append(f"Chart {ser.name}: " + ", ".join(f"{c}={v}" for c, v in zip(cats, ser.values)))
                    except Exception:
                        pass
            if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
                texts.append("Speaker notes: " + s.notes_slide.notes_text_frame.text.strip())
            texts = [t for t in texts if "Fictitious data for testing" not in t]
            sections.append({"label": f"Slide {i}: {slide_title or (texts[0][:60] if texts else '')}", "text": "\n".join(texts)})

    elif ext == "pdf":
        r = PdfReader(str(path))
        if r.metadata and r.metadata.title:
            title = r.metadata.title
        for i, pg in enumerate(r.pages, 1):
            t = pg.extract_text() or ""
            t = "\n".join(l for l in t.splitlines() if "Fictitious data for testing" not in l)
            sections.append({"label": f"Page {i}", "text": t})
        preface = sections[0]["text"] if sections else ""

    elif ext == "xlsx":
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        category = wb.properties.category or None
        title = wb.properties.title or title
        for ws in wb.worksheets:
            values = [list(r) for r in ws.iter_rows(values_only=True)]
            found = _find_table(values)
            if not found:
                txt = "\n".join(" | ".join(str(c) for c in r if c not in (None, "")) for r in values if any(r))
                if txt.strip():
                    sections.append({"label": f"Sheet {ws.title}", "text": txt})
                continue
            hi, headers, data = found
            intro = " ".join(str(c) for r in values[:hi] for c in r if c not in (None, ""))
            tables.append({"sheet": ws.title, "headers": headers, "rows": data, "intro": intro})
            secs = _table_sections(f"Sheet {ws.title}", headers, data)
            if secs and intro:
                secs[0]["text"] = intro + "\n" + secs[0]["text"]
            sections.extend(secs)
        wb.close()

    elif ext == "csv":
        with open(path, newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
        headers, data = rows[0], rows[1:]
        conv = []
        for r in data:
            out = []
            for v in r:
                try:
                    out.append(float(v) if re.fullmatch(r"-?\d+(\.\d+)?", v) else v)
                except ValueError:
                    out.append(v)
            conv.append(out)
        tables.append({"sheet": "csv", "headers": headers, "rows": conv, "intro": ""})
        sections.extend(_table_sections("Rows", headers, conv))

    elif ext == "json":
        obj = json.load(open(path, encoding="utf-8"))
        meta = {k: v for k, v in obj.items() if not isinstance(v, (list, dict))} if isinstance(obj, dict) else {}
        records_key = next((k for k, v in obj.items() if isinstance(v, list) and v and isinstance(v[0], dict)), None) \
            if isinstance(obj, dict) else None
        records = obj[records_key] if records_key else (obj if isinstance(obj, list) else [obj])
        if meta:
            preface = "; ".join(f"{k}: {v}" for k, v in meta.items())

        def flat(o, prefix=""):
            lines = []
            if isinstance(o, dict):
                for k, v in o.items():
                    lines += flat(v, f"{prefix}{k}.")
            elif isinstance(o, list):
                if all(not isinstance(x, (dict, list)) for x in o):
                    lines.append(f"{prefix[:-1]}: {', '.join(str(x) for x in o)}")
                else:
                    for i, x in enumerate(o):
                        lines += flat(x, f"{prefix}{i}.")
            elif o not in (None, ""):
                lines.append(f"{prefix[:-1]}: {o}")
            return lines
        for i, rec in enumerate(records):
            label = rec.get("id") or rec.get("bid_id") or rec.get("meeting_id") or rec.get("account") or f"record {i + 1}" \
                if isinstance(rec, dict) else f"record {i + 1}"
            text = "\n".join(flat(rec))
            sections.append({"label": f"{records_key or 'record'}: {label}", "text": (preface + "\n" + text) if i == 0 and preface else text})

    else:  # txt, md
        text = path.read_text(encoding="utf-8")
        first = text.strip().splitlines()[0] if text.strip() else title
        if ext == "txt" and first.lower().startswith(("from:", "microsoft teams", "[pasted")):
            m = re.search(r"^Subject: (.+)$", text, re.M) or re.search(r"^Meeting: (.+)$", text, re.M)
            if m:
                title = m.group(1).strip()
        elif first and len(first) < 120:
            title = first.strip("# ").strip()
        paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        for i, p in enumerate(paras):
            sections.append({"label": f"Part {i + 1}", "text": p})

    alltext = preface + "\n" + "\n".join(s["text"] for s in sections[:3])
    return {"title": str(title), "format": ext, "sections": [s for s in sections if s["text"].strip()],
            "tables": tables, "classification": classification(path, alltext, category)}


def chunk_sections(doc_title, sections, max_chars=1400, overlap=150):
    """Merge small sections, split large ones. Each chunk keeps its section label for citations."""
    chunks = []
    for s in sections:
        text = s["text"].strip()
        if len(text) <= max_chars:
            pieces = [text]
        else:
            pieces, buf = [], ""
            for line in text.split("\n"):
                if len(buf) + len(line) + 1 > max_chars and buf:
                    pieces.append(buf)
                    buf = buf[-overlap:] + "\n" + line
                else:
                    buf = (buf + "\n" + line) if buf else line
            if buf:
                pieces.append(buf)
        for p in pieces:
            chunks.append({"section": s["label"], "text": p})
    # merge tiny neighbours (e.g. one-line paragraphs of an email)
    merged = []
    for c in chunks:
        if merged and len(merged[-1]["text"]) + len(c["text"]) < 600 and merged[-1]["section"].split(" ")[0] == c["section"].split(" ")[0]:
            merged[-1]["text"] += "\n" + c["text"]
            merged[-1]["section"] = merged[-1]["section"].split(" - ")[0] + " - " + c["section"]
        else:
            merged.append(dict(c))
    return merged
